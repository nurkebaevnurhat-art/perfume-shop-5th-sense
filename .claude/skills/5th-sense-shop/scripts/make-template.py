# Шаблон каталога assets/files/5th-sense-catalog.xlsx из data.js (одна строка = один аромат, цена за 1 мл).
# Запуск из корня репозитория (нужен openpyxl):
#   node -e "const vm=require('vm'),fs=require('fs');const c={window:{FS:{}}};c.FS=c.window.FS;vm.createContext(c);for(const f of ['assets/js/config.js','assets/js/data.js'])vm.runInContext(fs.readFileSync(f,'utf8'),c);const F=c.FS;fs.writeFileSync('/tmp/catalog.json',JSON.stringify({families:F.families,products:F.defaultProducts,volumes:F.config.decants.volumes}))"
#   python3 .claude/skills/5th-sense-shop/scripts/make-template.py /tmp/catalog.json assets/files/5th-sense-catalog.xlsx
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

data = json.load(open(sys.argv[1]))
out = sys.argv[2]
VOLS = data['volumes']
TRACK = data.get('trackStock', False)  # остатки не ведём — колонки «Остаток» нет
fam = {f['id']: f['name'] for f in data['families']}
GENDER = {'men': 'Мужской', 'women': 'Женский', 'unisex': 'Унисекс'}
yn = lambda v: 'да' if v else ''

# (заголовок, вид: req|opt|calc, ширина, подсказка)
COLS = [
    ('Артикул', 'opt', 22, 'Необязательно. Если пусто, аромат определяется по бренду, названию и концентрации.'),
    ('Бренд', 'req', 24, 'Например: Tom Ford'),
    ('Название', 'req', 26, 'Например: Oud Wood'),
    ('Концентрация', 'opt', 14, 'EDP, EDT, Parfum, Extrait или Cologne'),
    ('Пол', 'req', 12, 'Мужской, Женский или Унисекс'),
    ('Цена за 1 мл, ₸', 'req', 14, 'Главная цена аромата. Только число. Цены 5, 10, 15 и 20 мл посчитаются сами: цена за 1 мл × миллилитры'),
] + [(f'Цена {ml} мл', 'calc', 11, f'Считается сама: цена за 1 мл × {ml}. Не меняйте — при загрузке колонка не читается') for ml in VOLS] + [
    ('Объёмы, мл', 'opt', 14, 'Пусто = ' + ', '.join(map(str, VOLS)) + ' мл. Если аромат разливаете не во всех объёмах, перечислите через запятую, например 5, 10'),
] + ([('Остаток, шт.', 'req', 12, 'Сколько штук каждого объёма в наличии, одно число. Можно по объёмам: 5: 3, 10: 2. 0 — нет в наличии')] if TRACK else []) + [
    ('Семейство', 'opt', 14, 'Выберите из списка'),
    ('Верхние ноты', 'opt', 28, 'Через запятую: бергамот, розовый перец'),
    ('Ноты сердца', 'opt', 28, 'Через запятую'),
    ('Базовые ноты', 'opt', 28, 'Через запятую'),
    ('Кратко', 'opt', 40, 'Одна фраза для карточки товара'),
    ('Описание', 'opt', 60, 'Несколько предложений для страницы товара'),
    ('Нишевая', 'opt', 10, 'да или пусто'),
    ('Бестселлер', 'opt', 11, 'да или пусто'),
    ('Новинка', 'opt', 10, 'да или пусто'),
    ('Год', 'opt', 8, 'Год выпуска аромата, например 2015'),
    ('Страна', 'opt', 14, 'Страна парфюмерного дома'),
    ('Парфюмер', 'opt', 22, 'Автор аромата'),
    ('Фото', 'opt', 30, 'Имя файла фото, например tom-ford-oud-wood.jpg. Фото загружаются отдельно'),
    ('Полный флакон, мл', 'opt', 14, 'Объём полноразмерного флакона, который можно заказать у поставщика (например 100). Необязательно'),
    ('Скрыть с сайта', 'opt', 12, 'да — аромат не показывается покупателям (например, закончился надолго). Пусто — виден'),
]
idx = {c[0]: i + 1 for i, c in enumerate(COLS)}
L = lambda name: get_column_letter(idx[name])
PER = L('Цена за 1 мл, ₸')

F = 'Arial'
fills = {'req': PatternFill('solid', fgColor='E9DCC3'), 'opt': PatternFill('solid', fgColor='F4EFE6'), 'calc': PatternFill('solid', fgColor='EEEEEE')}
calc_fill = PatternFill('solid', fgColor='F6F6F6')
thin = Side(style='thin', color='C9B99B')

wb = Workbook()
ws = wb.active
ws.title = 'Товары'
for i, (name, kind, width, hint) in enumerate(COLS, 1):
    c = ws.cell(row=1, column=i, value=name + (' *' if kind == 'req' else ''))
    c.font = Font(name=F, bold=True, size=10, italic=kind == 'calc', color='777777' if kind == 'calc' else None)
    c.fill = fills[kind]
    c.alignment = Alignment(vertical='center', wrap_text=True)
    c.border = Border(bottom=thin)
    c.comment = Comment(hint, '5th SENSE')
    ws.column_dimensions[get_column_letter(i)].width = width
ws.row_dimensions[1].height = 30
ws.freeze_panes = 'D2'

last = 1000
row = 2
for p in data['products']:
    n = p.get('notes') or {}
    vols = sorted(p['volumes'], key=lambda v: v['ml'])
    stocks = [v['stock'] for v in vols]
    mls = ', '.join(str(v['ml']) for v in vols)
    vals = {
        'Артикул': p['id'], 'Бренд': p['brand'], 'Название': p['name'], 'Концентрация': p.get('concentration', ''),
        'Пол': GENDER.get(p.get('gender'), ''), 'Цена за 1 мл, ₸': p.get('perMl') or round(vols[0]['price'] / vols[0]['ml']),
        'Объёмы, мл': '' if mls == ', '.join(map(str, VOLS)) else mls,
        'Остаток, шт.': stocks[0] if len(set(stocks)) == 1 else ', '.join(f"{v['ml']}: {v['stock']}" for v in vols),
        'Семейство': fam.get(p.get('family'), ''), 'Верхние ноты': ', '.join(n.get('top', [])),
        'Ноты сердца': ', '.join(n.get('heart', [])), 'Базовые ноты': ', '.join(n.get('base', [])),
        'Кратко': p.get('short', ''), 'Описание': p.get('description', ''),
        'Нишевая': yn(p.get('niche')), 'Бестселлер': yn(p.get('bestseller')), 'Новинка': yn(p.get('isNew')),
        'Год': p.get('year', ''), 'Страна': p.get('country', ''), 'Парфюмер': p.get('perfumer', ''),
        'Фото': (p.get('image') or '').split('/')[-1], 'Полный флакон, мл': p.get('fullMl', ''),
        'Скрыть с сайта': yn(p.get('hidden')),
    }
    for name, kind, *_ in COLS:
        if kind == 'calc':
            continue
        val = vals.get(name, '')
        c = ws.cell(row=row, column=idx[name], value=val if val != '' else None)
        c.font = Font(name=F, size=10, bold=name == 'Цена за 1 мл, ₸')
        c.alignment = Alignment(vertical='top', wrap_text=name in ('Кратко', 'Описание'))
    ws.cell(row=row, column=idx['Цена за 1 мл, ₸']).number_format = '#,##0'
    row += 1
# Формулы цен объёмов — на все строки, в том числе пустые для новых ароматов.
for r in range(2, last + 1):
    for ml in VOLS:
        c = ws.cell(row=r, column=idx[f'Цена {ml} мл'], value=f'=IF({PER}{r}="","",ROUND({PER}{r}*{ml},0))')
        c.number_format = '#,##0'
        c.font = Font(name=F, size=10, color='777777')
        c.fill = calc_fill

ls = wb.create_sheet('Списки')
lists = {'A': ('Пол', ['Мужской', 'Женский', 'Унисекс']), 'B': ('Концентрация', ['EDP', 'EDT', 'Parfum', 'Extrait', 'Cologne']),
         'C': ('Семейство', [f['name'] for f in data['families']]), 'D': ('Да/нет', ['да'])}
for col, (title, items) in lists.items():
    ls[f'{col}1'] = title
    ls[f'{col}1'].font = Font(name=F, bold=True)
    for i, it in enumerate(items, 2):
        ls[f'{col}{i}'] = it
        ls[f'{col}{i}'].font = Font(name=F)
    ls.column_dimensions[col].width = 16
def dv(formula, rng, title):
    d = DataValidation(type='list', formula1=formula, allow_blank=True, showErrorMessage=True, errorTitle=title, error='Выберите значение из списка')
    ws.add_data_validation(d); d.add(rng)
dv("='Списки'!$A$2:$A$4", f"{L('Пол')}2:{L('Пол')}{last}", 'Пол')
dv("='Списки'!$B$2:$B$6", f"{L('Концентрация')}2:{L('Концентрация')}{last}", 'Концентрация')
dv(f"='Списки'!$C$2:$C${1+len(data['families'])}", f"{L('Семейство')}2:{L('Семейство')}{last}", 'Семейство')
dv("='Списки'!$D$2:$D$2", f"{L('Нишевая')}2:{L('Новинка')}{last}", 'Да/нет')
dv("='Списки'!$D$2:$D$2", f"{L('Скрыть с сайта')}2:{L('Скрыть с сайта')}{last}", 'Да/нет')
d = DataValidation(type='decimal', operator='greaterThan', formula1='0', allow_blank=True, showErrorMessage=True, errorTitle='Цена за 1 мл', error='Введите число больше нуля, без пробелов и букв')
ws.add_data_validation(d); d.add(f'{PER}2:{PER}{last}')

ins = wb.create_sheet('Инструкция', 0)
ins.column_dimensions['A'].width = 110
ex = data['products'][0]
per = ex.get('perMl')
sp = lambda x: f'{x:,}'.replace(',', ' ')
lines = [
    ('Каталог 5th SENSE: как заполнять', True), ('', False),
    ('Одна строка = один аромат', True),
    ('Главное число — «Цена за 1 мл». Цены распива сайт считает сам: цена за 1 мл × миллилитры, ровно, без скидок.', False),
    (f'Например, {ex["brand"]} {ex["name"]}: цена за 1 мл {sp(per)} ₸ → ' + ', '.join(f'{ml} мл = {sp(per*ml)} ₸' for ml in VOLS) + '.', False),
    ('Серые колонки «Цена 5 мл» … «Цена 20 мл» считаются сами — по ним удобно проверить цены. Их менять не нужно.', False),
    ('', False),
    ('Как поменять цену', True),
    ('1. Найдите строку аромата (Ctrl+F или поиск в приложении).', False),
    ('2. Впишите новую цену за 1 мл. Серые колонки сразу покажут новые цены объёмов.', False),
    ('3. Сохраните файл и загрузите его в админке: «Товары» → «Загрузить таблицу» → проверьте, что изменится → «Применить».', False),
    ('Поднять цены сразу всем ароматам или одному бренду на N % проще в админке: «Товары» → «Процент на бренд или все».', False),
    ('', False),
    ('Колонки', True),
    ('Колонки с * обязательны: бренд, название, пол, цена за 1 мл' + (', остаток.' if TRACK else '.'), False),
    ('«Объёмы, мл»: пусто — аромат продаётся в объёмах ' + ', '.join(map(str, VOLS)) + ' мл. Если только в некоторых, перечислите: 5, 10.', False),
    (('«Остаток, шт.»: одно число — столько штук каждого объёма. Можно по объёмам: 5: 3, 10: 2. 0 — нет в наличии.' if TRACK
      else 'Остатки не ведём: любой аромат всегда можно заказать, наличие подтверждает менеджер.'), False),
    ('«Скрыть с сайта»: «да» — аромат не виден покупателям (например, закончился надолго). Пусто — виден.', False),
    ('Полноразмерный флакон отдельно не нужен: он есть у каждого аромата, цену подтверждает менеджер.', False),
    ('Ноты пишите через запятую. Пол, концентрация, семейство и «да» выбираются из списка в ячейке.', False),
    ('Наведите на заголовок колонки — появится подсказка.', False),
    ('', False),
    ('Что уже в таблице', True),
    (f'В листе «Товары» уже стоят {len(data["products"])} ароматов, которые сейчас продаются на сайте. Новые добавляйте ниже.', False),
    ('Аромат, которого нет в таблице, при загрузке пропадёт с сайта. Чтобы временно скрыть аромат, поставьте «да» в «Скрыть с сайта».', False),
    ('', False),
    ('Фото', True),
    ('В колонке «Фото» укажите имя файла (например, tom-ford-oud-wood.jpg) и загрузите фото в админке кнопкой «Загрузить фото».', False),
    ('Лучше всего квадратные фото на белом фоне, от 1000×1000 пикселей. Без фото на сайте показывается рисунок флакона.', False),
]
for i, (t, bold) in enumerate(lines, 1):
    c = ins.cell(row=i, column=1, value=t)
    c.font = Font(name=F, bold=bold, size=12 if bold else 10)
wb.active = 1
wb.save(out)
print('rows', row - 2)
